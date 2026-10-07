#!/usr/bin/env python3
"""IRP 2025 cost set for the pathway sensitivity (user, 7 Oct 2026). Writes costset_irp2025.json.

  python3 -I costset_irp2025.py <IRP 2025 assumptions workbook .xlsx> [out.json]

Recomputed from the workbook every time; never edit the JSON by hand.
Source: IRP 2025 (Gazette 53596) assumptions workbook, sheets 'Learning Rates' (capex by year, base case)
and 'New Tech Properties' (CCGT 2x1 9F.05 fixed and variable O&M, heat rate, gas price), January 2024
rands. Wind and solar fixed O&M are not in the workbook: they come from the IPP Office's REIPPPP
historical data note of 7 Nov 2025, the source of the workbook's BW6 lines (BW6 wind compliant 366.95,
BW6 solar preferred 269.04 R/kW/yr, January 2024 terms).
Escalated to 2026 rands by SA headline CPI, January 2024 to January 2026: 1.032 x 1.035 (Stats SA,
3.2% y/y January 2025, 3.5% y/y January 2026).
Not in the workbook, so the model's own values stay: rooftop, vanadium flow, iron-air, pumped hydro,
lives (except CCGT, 30 years in both), discount rate, battery fixed O&M (2.5% of capex, about the BW1
BESS ratio of 678 on 27,888).
The workbook has three lithium base-case columns with no source label; the first is used.
"""
import json, sys
import openpyxl

CPI = 1.032 * 1.035
COLS = {            # 0-based column in 'Learning Rates', base case
    'wind': 30,     # Onshore Wind (BW6) Base Case
    'pv': 33,       # Solar PV (BW6 Preferred) Base Case, single-axis tracking
    'batt': 20,     # Lithium Ion Base Case (100 MW/4 hr), first of three
    'ccgt': 5,      # Multi-shaft CTCC (2x1 925 MW, 9F.05)
    'offshore': 12, # Offshore Wind Floating Base Case (the model's offshore is floating)
}
wb = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
rows = list(wb['Learning Rates'].iter_rows(values_only=True))
hdr = {j: (str(rows[1][j]) + ' / ' + str(rows[2][j])).replace('\n', ' ') for j in COLS.values()}
capex = {}
for t, j in COLS.items():
    capex[t] = {}
    for r in rows[3:]:
        y = r[0]
        if isinstance(y, (int, float)) and 2026 <= y <= 2040:
            capex[t][str(int(y))] = round(r[j] * CPI, 1)
    assert len(capex[t]) == 15, (t, len(capex[t]))

ntp = list(wb['New Tech Properties'].iter_rows(values_only=True))
def row(label):
    for r in ntp:
        if r[1] and str(r[1]).strip().startswith(label): return r   # labels sit in column 1
    raise KeyError(label)
# The 2x1 9F.05 CCGT column in 'New Tech Properties', found by its label (column 1 holds the row labels).
CC = next(k for k, c in enumerate(ntp[3]) if isinstance(c, str) and c.strip().startswith('2x1 9F.05'))
fom_ccgt, vom_ccgt = row('Fixed O&M')[CC], row('Variable O&M')[CC]
heat, fuel = row('Heat rate')[CC], row('First year, ZAR/GJ')[CC]

out = {
  'meta': {'source': 'IRP 2025 assumptions workbook (Learning Rates, New Tech Properties); IPP Office REIPPPP historical data note, 7 Nov 2025 (wind and solar fixed O&M)',
           'basis': 'January 2024 rands x %.4f (SA CPI Jan 2024 to Jan 2026) = 2026 rands' % CPI,
           'columns': hdr, 'built_by': 'harness/pathway/costset_irp2025.py'},
  'capex': capex,                                            # R/kW, 2026 rands, by build year
  'fom': {'wind': round(366.95 * CPI, 1), 'pv': round(269.04 * CPI, 1),
          'ccgt': round(fom_ccgt * CPI, 1), 'offshore': None},   # offshore: see below
  'state': {'costCcgt': round(heat / 1000 * fuel * CPI, 1),     # R/MWh fuel: heat rate GJ/MWh x R/GJ
            'vomCcgt': round(vom_ccgt * CPI, 1)},
}
# Offshore floating fixed O&M sits in the non-fossil block of 'New Tech Properties'.
for r in ntp:
    cells = list(r)
    for k, c in enumerate(cells):
        if isinstance(c, str) and c.strip() == 'OSW, floating':
            col = k
            break
    else:
        continue
    break
for r in ntp:
    if any(isinstance(c, str) and c.strip().startswith('Fixed O&M') for c in r[col-3:col]):
        out['fom']['offshore'] = round(r[col] * CPI, 1); break
assert out['fom']['offshore'], 'offshore fixed O&M not found'
json.dump(out, open(sys.argv[2] if len(sys.argv) > 2 else 'costset_irp2025.json', 'w'), indent=1)
print(json.dumps({k: (v if k != 'capex' else {t: [c['2026'], c['2040']] for t, c in v.items()}) for k, v in out.items() if k != 'meta'}))
